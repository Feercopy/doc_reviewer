from copy import deepcopy
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.services.new_summary_traction import needs_verified_revenue_total, with_traction_totals


def test_v8_missing_traction_cells_use_em_dash_without_touching_numbers_or_mismatch():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2026", "Total"], "rows": [
            {"label": "DTB Uplift (Cumul)", "values": ["-", "—"]},
            {"label": "Revenue from DTB", "values": ["10", "невозможно извлечь данные"]},
            {"label": "Revenue non-DTB", "values": ["–", "no data in the document"]},
            {"label": "Total Revenue", "values": ["12", "12"], "mismatch_periods": ["2026"]},
        ],
    }}
    result = with_traction_totals(payload)["traction_summary"]["rows"]
    assert [row["values"] for row in result] == [["—", "—"], ["10", "—"], ["—", "—"], ["12", "12"]]
    assert result[-1]["mismatch_periods"] == ["2026"]


def test_new_traction_total_revenue_sums_complete_annual_values_only():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total 2026–2030"], "rows": [
            {"label": "DTB Uplift (Cumul)", "values": ["1%", "2%", "—"]},
            {"label": "Revenue from DTB", "values": ["1", "2", "—"]},
            {"label": "Revenue non-DTB", "values": ["3", "4", "—"]},
            {"label": "Total Revenue", "values": ["1 200,5", "2 300,5", "—"]},
        ],
    }}
    source = [{"metadata": {"rows": [
        ["ToBe P&L", "2027", "2028", "2027-28 total"], ["Revenue", "1 200,5", "2 300,5", "9 999"],
    ]}}]
    result = with_traction_totals(payload, source_blocks=source, strict_total_source=True)
    table = result["traction_summary"]
    assert table["total_source_policy"] == "same_table_verified_total"
    assert table["total_revenue_provenance"] == {"method": "annual_sum"}
    assert table["periods"] == ["2027", "2028", "Total 2027–2028"]
    assert [row["values"][-1] for row in table["rows"]] == ["—", "—", "—", "3501"]
    assert not needs_verified_revenue_total(result)
    assert with_traction_totals(result, source_blocks=source) == result
    assert payload["traction_summary"]["periods"][-1] == "Total 2026–2030"


def test_new_traction_rejects_unverified_total_even_when_populated():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "DTB Uplift (Cumul)", "values": ["—", "—", "—"]},
            {"label": "Revenue from DTB", "values": ["—", "—", "—"]},
            {"label": "Revenue non-DTB", "values": ["—", "—", "—"]},
            {"label": "Total Revenue", "values": ["10", "20", "99"]},
        ],
    }}
    result = with_traction_totals(payload, strict_total_source=True)
    assert result["traction_summary"]["rows"][-1]["values"][-1] == "30"
    assert result["traction_summary"]["total_revenue_provenance"] == {"method": "annual_sum"}
    assert payload["traction_summary"]["rows"][-1]["values"][-1] == "99"
    for annual in (["10", "—"], ["1,234", "2"], ["10 (H2 2027)", "20"]):
        incomplete = deepcopy(payload)
        incomplete["traction_summary"]["rows"][-1]["values"] = [*annual, "99"]
        result = with_traction_totals(incomplete, strict_total_source=True)
        assert result["traction_summary"]["rows"][-1]["values"][-1] == "—"
        assert result["traction_summary"]["total_revenue_provenance"] == {"method": "unavailable"}


def test_new_traction_accepts_total_only_from_matching_source_row_and_horizon():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "DTB Uplift (Cumul)", "values": ["—", "—", "—"]},
            {"label": "Revenue from DTB", "values": ["—", "—", "—"]},
            {"label": "Revenue non-DTB", "values": ["—", "—", "—"]},
            {"label": "Total Revenue", "values": ["10", "20", "999"], "mismatch_periods": ["2027", "Total"]},
        ],
    }}
    correct = [{"metric": "revenue", "source_block_id": "financial-model-4", "source_page": 4,
                "periods": ["2027", "2028", "Total 2027–2028"],
                "rows": [{"label": "Revenue", "values": ["10", "20", "35"]}]}]
    table = with_traction_totals(payload, source_tables=correct, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"] == ["10", "20", "35"]
    assert table["rows"][-1]["mismatch_periods"] == ["2027"]
    assert table["total_revenue_provenance"] == {
        "method": "same_source_table", "source_block_id": "financial-model-4", "source_page": 4,
    }

    wrong_years = deepcopy(correct)
    wrong_years[0]["rows"][0]["values"] = ["11", "20", "35"]
    wrong_horizon = deepcopy(correct)
    wrong_horizon[0]["periods"][-1] = "Total 2026–2028"
    for source in (wrong_years, wrong_horizon):
        table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
        assert table["rows"][-1]["values"][-1] == "30"
        assert table["total_revenue_provenance"] == {"method": "annual_sum"}


def test_new_traction_rejects_ambiguous_grouped_source_numbers():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Total Revenue", "values": ["10", "20", "999"]},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["10,000", "20,000", "30,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "30"
    assert table["total_revenue_provenance"] == {"method": "annual_sum"}

    source[0]["rows"][0]["values"] = ["10", "20", "30,000"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "30"

    signed = deepcopy(payload)
    signed["traction_summary"]["rows"][0]["values"] = ["-1", "-2", "999"]
    for grouped in ("(1,000)", "−1,000", "-1, 000", "-1.000"):
        source[0]["rows"][0]["values"] = [grouped, "-2", "-3,000"]
        table = with_traction_totals(signed, source_tables=source, strict_total_source=True)["traction_summary"]
        assert table["rows"][0]["values"][-1] == "-3"
        assert table["total_revenue_provenance"] == {"method": "annual_sum"}

    source[0]["rows"][0]["values"] = ["10", "20", "3,000,5"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "30"
    assert table["total_revenue_provenance"] == {"method": "annual_sum"}


def test_new_traction_preserves_exact_grouped_values_from_one_source_row():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["1,000", "2,000", "999"]}],
    }}
    source = [{"metric": "revenue", "source_block_id": "faq-3-table", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,000", "2,000", "3,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "3,000"
    assert table["total_revenue_provenance"] == {"method": "same_source_table", "source_block_id": "faq-3-table"}

    payload["traction_summary"]["rows"][0]["values"][0] = "1"
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "—"
    assert table["total_revenue_provenance"] == {"method": "unavailable"}


def test_new_traction_allows_sub_thousand_cell_in_grouped_source_row():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["500", "1,000", "999"]}],
    }}
    source = [{"metric": "revenue", "source_block_id": "financial-model-5",
               "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["500", "1,000", "1,500"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "1,500"
    assert table["total_revenue_provenance"] == {"method": "same_source_table", "source_block_id": "financial-model-5"}


def test_new_traction_matches_zero_across_grouped_decimal_notation():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["0", "0", "999"]}],
    }}
    source = [{"metric": "revenue", "source_block_id": "financial-model-zero",
               "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["0,000", "0,000", "0,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "0,000"
    assert table["total_revenue_provenance"] == {"method": "same_source_table", "source_block_id": "financial-model-zero"}

    source[0]["rows"][0]["values"] = ["0", "0", "0,000"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "0,000"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"

    duplicate = deepcopy(source[0])
    duplicate["rows"][0]["values"][-1] = "0"
    table = with_traction_totals(payload, source_tables=[source[0], duplicate], strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "0,000"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"


def test_new_traction_preserves_repeated_grouped_total_from_matching_row():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["400,000", "600,000", "999"]}],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["400,000", "600,000", "1,000,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "1,000,000"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"

    payload["traction_summary"]["rows"][0]["values"] = ["1,000,000", "2,000,000", "999"]
    source[0]["rows"][0]["values"] = ["1,000,000", "2,000,000", "3,000,000"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "3,000,000"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"


def test_new_traction_accepts_grouping_with_decimal_fraction_from_source():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["1,000.50", "2,000.25", "999"]}],
    }}
    source = [{"metric": "revenue", "source_block_id": "financial-model-decimal",
               "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,000.50", "2,000.25", "3,000.75"]}]}]
    for first, second, total in (("1,000.50", "2,000.25", "3,000.75"),
                                 ("1.000,50", "2.000,25", "3.000,75")):
        payload["traction_summary"]["rows"][0]["values"] = [first, second, "999"]
        source[0]["rows"][0]["values"] = [first, second, total]
        table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
        assert table["rows"][0]["values"][-1] == total
        assert table["total_revenue_provenance"]["source_block_id"] == "financial-model-decimal"

    payload["traction_summary"]["rows"][0]["values"] = ["1,000", "2,000.00", "999"]
    source[0]["rows"][0]["values"] = ["1,000", "2,000.00", "3,000"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "3,000"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"


def test_new_traction_accepts_equal_unambiguous_duplicate_source_totals():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["10", "20", "999"]}],
    }}
    sources = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
                "rows": [{"label": "Revenue", "values": ["10", "20", total]}]}
               for total in ("35", "35.0")]
    table = with_traction_totals(payload, source_tables=sources, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "35"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"


def test_new_traction_uses_decimal_role_from_source_row():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["0.5", "1.25", "999"]}],
    }}
    source = [{"metric": "revenue", "source_block_id": "decimal-source",
               "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["0.500", "1.25", "1.750"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "1.750"
    assert table["total_revenue_provenance"]["source_block_id"] == "decimal-source"

    duplicate = deepcopy(source[0])
    duplicate["rows"][0]["values"][-1] = "1.75"
    table = with_traction_totals(payload, source_tables=[*source, duplicate], strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "1.750"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"

    payload["traction_summary"]["rows"][0]["values"] = ["0,500", "1,25", "999"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "1.750"
    assert table["total_revenue_provenance"]["method"] == "same_source_table"


def test_new_traction_reconciles_consistently_grouped_component_totals():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["500", "500", "1,000"]},
            {"label": "Revenue non-DTB", "values": ["1,000", "1,000", "2,000"]},
            {"label": "Total Revenue", "values": ["1,000", "2,000", "999"], "mismatch_periods": []},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,000", "2,000", "3,500"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "3,500"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    payload["traction_summary"]["rows"][0]["values"][-1] = "1 000"
    payload["traction_summary"]["rows"][1]["values"][-1] = "2 000"
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]


def test_new_traction_does_not_claim_mismatch_for_ambiguous_decimal_total():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["1.0", "0.5", "1.5"]},
            {"label": "Revenue non-DTB", "values": ["1.0", "1.0", "2.0"]},
            {"label": "Total Revenue", "values": ["1,500", "2,000", "999"], "mismatch_periods": ["Total"]},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,500", "2,000", "3,500"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "3,500"
    assert table["rows"][-1]["mismatch_periods"] == []


def test_new_traction_uses_one_interpretation_for_shared_separator():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["0", "0", "1,000"]},
            {"label": "Revenue non-DTB", "values": ["0", "0", "2,000"]},
            {"label": "Total Revenue", "values": ["1", "1", "999"], "mismatch_periods": []},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1", "1", "1002"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "1002"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    payload["traction_summary"]["rows"][0]["values"][-1] = "1000.50"
    payload["traction_summary"]["rows"][1]["values"][-1] = "2.000"
    source[0]["rows"][0]["values"][-1] = "3000.50"
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    payload["traction_summary"]["rows"][0]["values"][-1] = "1,000.50"
    payload["traction_summary"]["rows"][1]["values"][-1] = "2.000"
    source[0]["rows"][0]["values"][-1] = "3,000.50"
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    payload["traction_summary"]["rows"][0]["values"][-1] = "1,000,000"
    payload["traction_summary"]["rows"][1]["values"][-1] = "500,000"
    source[0]["rows"][0]["values"][-1] = "1,000,500"
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]


def test_zero_annual_does_not_establish_grouping_for_ambiguous_total():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["0", "20", "999"]}],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["0,000", "20", "3,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][0]["values"][-1] == "20"
    assert table["total_revenue_provenance"]["method"] == "annual_sum"


def test_new_traction_recomputes_total_mismatch_after_replacement():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["5", "5", "10"]},
            {"label": "Revenue non-DTB", "values": ["5", "15", "20"]},
            {"label": "Total Revenue", "values": ["10", "20", "999"], "mismatch_periods": ["Total"]},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["10", "20", "35"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "35"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    table = with_traction_totals(payload, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "30"
    assert table["rows"][-1]["mismatch_periods"] == []


def test_new_traction_reconciles_mismatch_when_verified_total_is_unchanged():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["200,000", "200,000", "400,000"]},
            {"label": "Revenue non-DTB", "values": ["250,000", "250,000", "500,000"]},
            {"label": "Total Revenue", "values": ["450,000", "450,000", "1,000,000"], "mismatch_periods": []},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["450,000", "450,000", "1,000,000"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "1,000,000"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]

    payload["traction_summary"]["rows"][1]["values"][-1] = "600,000"
    payload["traction_summary"]["rows"][-1]["mismatch_periods"] = ["Total"]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == []


def test_new_traction_reconciles_using_resolved_verified_total():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["1", "1", "1 000"]},
            {"label": "Revenue non-DTB", "values": ["1", "1", "2 500"]},
            {"label": "Total Revenue", "values": ["1,25", "2,25", "3,500"], "mismatch_periods": ["Total"]},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,25", "2,25", "3,500"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "3,500"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]


def test_verified_total_does_not_conflict_with_component_number_locale():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["5", "5", "1.000,00"]},
            {"label": "Revenue non-DTB", "values": ["5", "5", "2.000,00"]},
            {"label": "Total Revenue", "values": ["10", "20", "3000.50"], "mismatch_periods": []},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["10", "20", "3000.50"]}]}]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]


def test_verified_duplicate_total_is_not_source_order_dependent():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2027", "2028", "Total"], "rows": [
            {"label": "Revenue from DTB", "values": ["1", "1", "1.5"]},
            {"label": "Revenue non-DTB", "values": ["1", "1", "2"]},
            {"label": "Total Revenue", "values": ["1,500", "2,000", "999"], "mismatch_periods": []},
        ],
    }}
    source = [{"metric": "revenue", "periods": ["2027", "2028", "Total 2027–2028"],
               "rows": [{"label": "Revenue", "values": ["1,500", "2,000", total]}]}
              for total in ("3,500", "3500")]
    table = with_traction_totals(payload, source_tables=source, strict_total_source=True)["traction_summary"]
    assert table["rows"][-1]["values"][-1] == "3500"
    assert table["rows"][-1]["mismatch_periods"] == ["Total"]


def test_saved_same_table_annual_sum_report_is_not_recomputed():
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "total_source_policy": "same_table_annual_sum",
        "periods": ["2027", "2028", "Total"],
        "rows": [{"label": "Total Revenue", "values": ["10", "20", "99"]}],
    }}
    assert with_traction_totals(payload) == payload


def test_v7_single_table_keeps_four_rows_and_mismatch_without_legacy_rewrite():
    table = {"periods": ["2026", "Total"], "rows": [
        {"label": "DTB Uplift (Cumul)", "values": ["2%", "—"]},
        {"label": "Revenue from DTB", "values": ["10", "10"]},
        {"label": "Revenue non-DTB", "values": ["20", "20"]},
        {"label": "Total Revenue", "values": ["35", "35"], "mismatch_periods": ["2026", "Total"]},
    ]}
    payload = {"schema_version": "new-summary-v7", "language": "ru", "traction_summary": table}

    result = with_traction_totals(payload)

    assert result["traction_summary"] == table
    assert result is not payload


def test_historical_traction_break_marker_is_cleaned_without_changing_total():
    payload = {"language": "ru", "traction_summary": {"tables": [{
        "metric": "revenue", "metric_label": "Выручка (инкр.)",
        "periods": ["2026", "2026-31<br>total"],
        "rows": [{"label": "Revenue uplift", "values": ["10", "15 918"]}],
    }]}}
    result = with_traction_totals(payload)["traction_summary"]["tables"][0]
    assert result["periods"] == ["2026", "2026-31 total"]
    assert result["rows"][0]["values"] == ["10", "15 918"]
    assert payload["traction_summary"]["tables"][0]["periods"][-1] == "2026-31<br>total"
from app.services.new_summary_exports import (
    NewSummaryExportProvenance, _build_docx, _build_pdf, _confirmed_first,
    _gate2_hypothesis_heading, _labels, _validation_rationale,
)
from app.services.new_summaries import with_verified_source_totals
from app.schemas.analyses import NewSummaryRead, NewSummaryVariantRead
from uuid import uuid4


def _payload(language: str = "ru") -> dict:
    return {
        "language": language,
        "traction_summary": {"tables": [
            {"metric": "revenue", "metric_label": "Revenue", "periods": ["2026", "2027"],
             "rows": [{"label": "Incremental revenue", "values": ["10", "20"]}]},
            {"metric": "dtb", "metric_label": "DTB", "periods": ["CY26", "CY27"],
             "rows": [{"label": "DTB uplift", "values": ["1%", "2%"]}]},
        ]},
    }


def test_missing_total_is_explicit_and_never_calculated():
    payload = _payload()
    original = deepcopy(payload)
    tables = with_traction_totals(payload)["traction_summary"]["tables"]
    assert tables[0]["periods"] == ["2026", "2027", "Total 2026–2027"]
    assert tables[0]["rows"][0]["values"] == ["10", "20", "невозможно извлечь данные"]
    assert tables[1]["periods"][-1] == "Total"
    assert payload == original


def test_verified_source_without_total_uses_document_absence_reason():
    payload = _payload()
    source = [{"metric": "revenue", "periods": ["2026", "2027"]}]
    tables = with_traction_totals(payload, source_tables=source)["traction_summary"]["tables"]
    assert tables[0]["rows"][0]["values"][-1] == "отсутствуют данные в документе защиты"
    assert tables[1]["rows"][0]["values"][-1] == "невозможно извлечь данные"


def test_existing_total_values_are_preserved_in_both_languages():
    for language in ("ru", "en"):
        payload = _payload(language)
        table = payload["traction_summary"]["tables"][0]
        table["periods"].append("2026–2027 total")
        table["rows"][0]["values"].append("17")
        result = with_traction_totals(payload)["traction_summary"]["tables"]
        assert result[0] == table
        assert len(result[0]["periods"]) == len(result[0]["rows"][0]["values"])
        assert result[1]["rows"][0]["values"][-1] == (
            "невозможно извлечь данные" if language == "ru" else "could not extract data"
        )


def test_existing_total_moves_to_last_column_with_matching_value():
    payload = _payload()
    table = payload["traction_summary"]["tables"][0]
    table["periods"] = ["2026", "Total", "2027"]
    table["rows"][0]["values"] = ["10", "30", "20"]
    result = with_traction_totals(payload)["traction_summary"]["tables"][0]
    assert result["periods"] == ["2026", "2027", "Total"]
    assert result["rows"][0]["values"] == ["10", "20", "30"]


def test_saved_summary_placeholder_is_enriched_only_by_verified_source_total():
    payload = _payload()
    table = payload["traction_summary"]["tables"][0]
    table["periods"].append("Total 2026–2027")
    table["rows"][0]["values"].append("невозможно извлечь данные")
    source = [{"id": "faq6", "metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"],
        ["Revenue", "10", "20", "30"],
    ]}}]
    enriched = with_traction_totals(payload, source_blocks=source)
    assert enriched["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "30"
    assert payload["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "невозможно извлечь данные"
    source[0]["metadata"]["rows"][1][2] = "21"
    assert with_traction_totals(payload, source_blocks=source)["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "невозможно извлечь данные"


def test_old_bilingual_report_reads_verified_total_without_mutating_saved_payload(monkeypatch):
    import app.services.new_summaries as summaries

    source = [{"id": "faq6", "metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"],
        ["Revenue", "10", "20", "30"],
    ]}}]
    monkeypatch.setattr(summaries, "verified_table_blocks", lambda _document: source)
    ru = _payload("ru")
    en = _payload("en")
    report = NewSummaryRead(
        analysis_id=uuid4(), ru=NewSummaryVariantRead(status="completed", payload=ru),
        en=NewSummaryVariantRead(status="completed", payload=en),
    )
    updated = with_verified_source_totals(report, object())
    for language in ("ru", "en"):
        assert getattr(updated, language).payload["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "30"
        assert getattr(report, language).payload["traction_summary"]["tables"][0]["periods"] == ["2026", "2027"]


def test_gate2_hypothesis_heading_uses_actual_items_and_confirmed_first():
    item = {"id": "gate2_hypothesis_results", "detail": {"type": "solution_validation", "items": [
        {"text": "a", "verdict": "insufficient"}, {"text": "b", "verdict": "confirmed"},
        {"text": "c", "verdict": "insufficient"},
    ]}}
    assert _gate2_hypothesis_heading(item, _labels("ru")) == (
        "Результаты проверки гипотез из Gate: 1 гипотез из 3 подтверждены, 2 гипотез из 3 недостаточно подтверждены."
    )
    assert [entry["text"] for entry in _confirmed_first(item["detail"]["items"], "verdict")] == ["b", "a", "c"]


def test_v2_solution_checks_show_counts_and_source_rationale():
    item = {"id": "stream_review_1_solution_validation", "detail": {"type": "solution_validation", "items": [
        {"text": "Pilot", "verdict": "confirmed", "test": "Pilot with 20 users",
         "expected_result": "10 activations", "actual_result": "12 activations"},
        {"text": "Launch", "verdict": "insufficient"},
    ]}}
    assert _gate2_hypothesis_heading(item, _labels("ru"), new_format=True) == (
        "Подтвержденные решения: 1 проверка из 2 подтверждена, 1 проверка из 2 недостаточно подтверждена."
    )
    assert _validation_rationale(item["detail"]["items"][0], _labels("ru")) == (
        "Проверка: Pilot with 20 users; ожидали: 10 activations; получили: 12 activations."
    )


def test_v2_docx_omits_removed_sections_and_numbers_required_elements():
    item = {"id": "stream_review_1_solution_validation", "label": "Подтверждение решения", "status": "1/1",
            "detail": {"type": "solution_validation", "items": [{
                "text": "Пилот", "verdict": "confirmed", "test": "Пилот на 20 пользователях",
                "expected_result": "10 активаций", "actual_result": "12 активаций",
            }]}}
    content = {"schema_version": "new-summary-v2", "language": "ru", "title": "AI Summary Test",
               "stage": "Stream Review 1", "context": "Проверен пилот.",
               "traction_summary": {"tables": []}, "required_elements": [item],
               "document_quality_percent": 75, "other": ["Legacy-only note"], "critical_problems": []}
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    document = DocxDocument(BytesIO(_build_docx({"ru": content, "en": {**content, "language": "en"}}, provenance)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "1. Подтвержденные решения: 1 проверка из 1 подтверждена" in text
    assert "Проверка: Пилот на 20 пользователях" in text
    assert "Качество документа" not in text
    assert "Другие наблюдения" not in text
    assert "Legacy-only note" not in text


def test_v3_problem_issue_is_bold_in_exports_and_v2_string_still_renders():
    content = {
        "schema_version": "new-summary-v3", "language": "en", "title": "AI Summary Test",
        "stage": "Gate 2", "context": "A pilot was conducted.",
        "traction_summary": {"tables": []}, "required_elements": [],
        "critical_problems": [{
            "issue": "Revenue starts late.",
            "fact": "The current plan shows no revenue in the first year.",
        }],
    }
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    report = {"ru": {**content, "language": "ru"}, "en": content}
    document = DocxDocument(BytesIO(_build_docx(report, provenance)))
    problem = next(paragraph for paragraph in document.paragraphs if "Revenue starts late" in paragraph.text)
    assert problem.runs[0].text == "Revenue starts late."
    assert problem.runs[0].bold is True
    assert problem.runs[1].text == " The current plan shows no revenue in the first year."
    assert not problem.runs[1].bold
    assert "Revenue starts late." in PdfReader(BytesIO(_build_pdf(report, provenance))).pages[0].extract_text()

    legacy = {**content, "schema_version": "new-summary-v2", "critical_problems": ["Legacy plain-text problem."]}
    legacy_report = {"ru": {**legacy, "language": "ru"}, "en": legacy}
    legacy_document = DocxDocument(BytesIO(_build_docx(legacy_report, provenance)))
    assert any("Legacy plain-text problem." in paragraph.text for paragraph in legacy_document.paragraphs)


def test_legacy_single_revenue_table_without_metric_can_receive_verified_total():
    payload = {"language": "ru", "traction_summary": {
        "metric_label": "Revenue (incr)", "periods": ["2026", "2027", "Total"],
        "rows": [{"label": "Revenue", "values": ["10", "20", "Не смог получить данные"]}],
    }}
    source = [{"metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"], ["Revenue", "10", "20", "30"],
    ]}}]
    assert needs_verified_revenue_total(payload)
    assert with_traction_totals(payload, source_blocks=source)["traction_summary"]["rows"][0]["values"][-1] == "30"


def test_v8_verified_total_applies_dtb_revenue_rule_without_inventing_missing_dtb():
    source = [{"metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"], ["Revenue", "10", "20", "30"],
    ]}}]
    payload = {"schema_version": "new-summary-v8", "traction_summary": {
        "periods": ["2026", "2027", "Total"], "rows": [
            {"label": "DTB Uplift (Cumul)", "values": ["1%", "2%", "—"]},
            {"label": "Revenue from DTB", "values": ["—", "—", "—"]},
            {"label": "Revenue non-DTB", "values": ["—", "—", "—"]},
            {"label": "Total Revenue", "values": ["10", "20", "—"]},
        ],
    }}
    enriched = with_traction_totals(payload, source_blocks=source)["traction_summary"]["rows"]
    assert enriched[1]["values"] == ["10", "20", "30"]
    assert enriched[3]["values"] == ["10", "20", "30"]
    without_dtb = {**payload, "traction_summary": {
        **payload["traction_summary"],
        "rows": [{**payload["traction_summary"]["rows"][0], "values": ["—", "—", "—"]}, *payload["traction_summary"]["rows"][1:]],
    }}
    no_copy = with_traction_totals(without_dtb, source_blocks=source)["traction_summary"]["rows"]
    assert no_copy[1]["values"] == ["—", "—", "—"]
    assert no_copy[3]["values"] == ["10", "20", "30"]
